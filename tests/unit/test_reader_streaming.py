import io
from unittest.mock import patch

import pypdfium2 as pdfium
import pytest
from docx import Document
from openpyxl import Workbook
from PIL import Image

from rag_pre_ingest.models.document import Page
from rag_pre_ingest.pipeline.readers import read_document


def source(kind: str) -> tuple[bytes, int]:
    stream = io.BytesIO()
    if kind == "pdf":
        with pdfium.PdfDocument.new() as document:
            for _ in range(3):
                document.new_page(100, 200).close()
            document.save(stream)
        return stream.getvalue(), 3
    if kind == "docx":
        document = Document()
        document.add_paragraph("Text")
        document.save(stream)
        return stream.getvalue(), 1
    if kind == "xlsx":
        workbook = Workbook()
        workbook.create_sheet("second")
        workbook.save(stream)
        return stream.getvalue(), 2
    Image.new("RGB", (40, 40)).save(
        stream, format="TIFF", save_all=True, append_images=[Image.new("RGB", (20, 20))]
    )
    return stream.getvalue(), 2


@pytest.mark.parametrize("kind", ["pdf", "docx", "xlsx", "tiff"])
def test_pages_are_dispatched_in_order_and_keep_identity(kind: str) -> None:
    data, count = source(kind)
    notified: list[Page] = []
    result = read_document(data, kind, "file", "hash", "id", on_page=notified.append)
    assert [page.page_number for page in notified] == list(range(1, count + 1))
    assert all(notified[index] is result.pages[index] for index in range(count))


@pytest.mark.parametrize("kind", ["pdf", "docx", "xlsx", "tiff"])
def test_callback_error_preserves_pages_and_is_not_source_corruption(kind: str) -> None:
    data, count = source(kind)

    def broken(page: Page) -> None:
        raise RuntimeError("scheduler unavailable")

    result = read_document(data, kind, "file", "hash", "id", on_page=broken)
    assert len(result.pages) == count
    assert sum(w.code == "page_callback_failed" for w in result.warnings) == count
    assert not any(w.code == "corrupted_page" for w in result.warnings)


def test_pdf_callback_runs_before_next_page_is_read() -> None:
    data, _ = source("pdf")
    seen: list[int] = []
    original = pdfium.PdfDocument.__getitem__

    def get_page(document: pdfium.PdfDocument, index: int) -> pdfium.PdfPage:
        assert seen == list(range(1, index + 1))
        return original(document, index)

    with patch.object(pdfium.PdfDocument, "__getitem__", get_page):
        read_document(
            data, "pdf", "file", "hash", "id", on_page=lambda p: seen.append(p.page_number)
        )
    assert seen == [1, 2, 3]
