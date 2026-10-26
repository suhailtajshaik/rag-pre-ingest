import io

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from rag_pre_ingest.pipeline.readers import read_document
from rag_pre_ingest.pipeline.sanitizer import sanitize_document


def test_word_bullets_and_consecutive_numbers() -> None:
    document = Document()
    document.add_paragraph("Bullet one", style="List Bullet")
    document.add_paragraph("Bullet two", style="List Bullet")
    document.add_paragraph("First numbered", style="List Number")
    document.add_paragraph("Second numbered", style="List Number")
    document.add_paragraph("Ordinary prose")
    stream = io.BytesIO()
    document.save(stream)
    result = read_document(stream.getvalue(), "docx", "lists.docx", "hash", "id")
    assert [block.type for block in result.pages[0].blocks] == ["list"] * 4 + ["paragraph"]
    assert "- Bullet one" in result.markdown and "- Bullet two" in result.markdown
    assert "1. First numbered" in result.markdown and "2. Second numbered" in result.markdown
    assert result.pages[0].blocks[-1].text == "Ordinary prose"
    clean = sanitize_document(result)
    assert "1. First numbered" in clean.markdown and "2. Second numbered" in clean.markdown


def test_nested_list_style_retains_indent() -> None:
    document = Document()
    document.add_paragraph("Outer", style="List Bullet")
    document.add_paragraph("Inner", style="List Bullet 2")
    stream = io.BytesIO()
    document.save(stream)
    result = read_document(stream.getvalue(), "docx", "lists.docx", "hash", "id")
    assert result.pages[0].blocks[1].markdown == "  - Inner"
    assert result.pages[0].blocks[1].bbox[0] > result.pages[0].blocks[0].bbox[0]


def test_explicit_unknown_number_format_warns_without_fabrication() -> None:
    document = Document()
    paragraph = document.add_paragraph("Roman item")
    root = document.part.numbering_part.element
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), "99")
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    fmt = OxmlElement("w:numFmt")
    fmt.set(qn("w:val"), "lowerRoman")
    level.append(fmt)
    abstract.append(level)
    root.append(abstract)
    number = OxmlElement("w:num")
    number.set(qn("w:numId"), "99")
    link = OxmlElement("w:abstractNumId")
    link.set(qn("w:val"), "99")
    number.append(link)
    root.append(number)
    properties = paragraph._p.get_or_add_pPr().get_or_add_numPr()
    properties.get_or_add_numId().val = 99
    stream = io.BytesIO()
    document.save(stream)
    result = read_document(stream.getvalue(), "docx", "lists.docx", "hash", "id")
    assert result.pages[0].blocks[0].text == "Roman item"
    assert any(w.code == "lists_uncertain" for w in result.warnings)


def test_explicit_numbering_start_override() -> None:
    document = Document()
    paragraph = document.add_paragraph("Starts at five", style="List Number")
    root = document.part.numbering_part.element
    style_num = paragraph.style.element.pPr.numPr.numId.val
    number = next(
        node for node in root.findall(qn("w:num")) if node.get(qn("w:numId")) == str(style_num)
    )
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "5")
    override.append(start)
    number.append(override)
    stream = io.BytesIO()
    document.save(stream)
    result = read_document(stream.getvalue(), "docx", "lists.docx", "hash", "id")
    assert result.pages[0].blocks[0].markdown == "5. Starts at five"


def test_nested_list_indentation_survives_final_sanitizer() -> None:
    from rag_pre_ingest.pipeline.sanitizer import sanitize_document

    document = Document()
    document.add_paragraph("Outer", style="List Bullet")
    document.add_paragraph("Inner", style="List Bullet 2")
    stream = io.BytesIO()
    document.save(stream)
    result = sanitize_document(read_document(stream.getvalue(), "docx", "list.docx", "hash", "id"))
    assert "  - Inner" in result.markdown
