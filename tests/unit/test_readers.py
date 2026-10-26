import io
import time

import pypdfium2 as pdfium
import pytest
from docx import Document
from openpyxl import Workbook
from PIL import Image

from rag_pre_ingest.pipeline.readers import (
    normalize_text,
    page_deadline,
    read_document,
    render_page,
    render_table,
)


def read(data: bytes, kind: str):
    return read_document(data, kind, "document." + kind, "hash", "document")


def test_normalize_text() -> None:
    assert normalize_text("  ﬁ\u00a0test\x00 ") == "fi test"


def test_render_table() -> None:
    assert render_table([["A", "B"], ["x"]]) == "| A | B |\n| --- | --- |\n| x |  |"
    assert render_table([]) == ""


def test_word_read_and_preview() -> None:
    doc = Document()
    doc.add_heading("Title", 1)
    doc.add_paragraph("Hello world")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    table.cell(1, 0).text = "value"
    output = io.BytesIO()
    doc.save(output)
    result = read(output.getvalue(), "docx")
    assert "# Title" in result.markdown
    assert "Hello world" in result.markdown
    assert result.stats.tables.total == 1
    assert result.pages[0].blocks[-1].cells[-1].text == ""
    assert any(w.code == "synthetic_page" for w in result.warnings)
    assert render_page(output.getvalue(), "docx", 1).startswith(b"\x89PNG")


def test_excel_blanks_and_merges() -> None:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.append(["A", "B", "C"])
    sheet.append(["x", None, "z"])
    sheet.merge_cells("A3:B3")
    sheet["A3"] = "merged"
    output = io.BytesIO()
    book.save(output)
    result = read(output.getvalue(), "xlsx")
    assert "| x |  | z |" in result.markdown
    assert result.stats.tables.with_merged_cells == 1
    assert render_page(output.getvalue(), "xlsx", 1).startswith(b"\x89PNG")


def test_empty_pdf_is_warning() -> None:
    output = io.BytesIO()
    with pdfium.PdfDocument.new() as doc:
        doc.new_page(100, 200).close()
        doc.save(output)
    result = read(output.getvalue(), "pdf")
    assert result.pages[0].width == 100
    assert result.markdown == ""
    assert any(w.code == "empty_document" for w in result.warnings)
    assert render_page(output.getvalue(), "pdf", 1).startswith(b"\x89PNG")


def test_image_no_markers_in_markdown() -> None:
    output = io.BytesIO()
    Image.new("RGB", (40, 30)).save(output, format="PNG")
    result = read(output.getvalue(), "png")
    assert result.markdown == ""
    assert result.stats.images.total == 1
    assert result.pages[0].blocks[0].type == "image"
    assert render_page(output.getvalue(), "png", 1).startswith(b"\x89PNG")


def test_corruption_becomes_warning() -> None:
    result = read(b"bad bytes", "pdf")
    assert any(w.code == "corrupted_page" for w in result.warnings)


def test_deadline_and_cleanup() -> None:
    with pytest.raises(TimeoutError), page_deadline(0.005):
        time.sleep(0.02)
    with page_deadline(1):
        assert True


def test_pdf_text_coordinates() -> None:
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 200 200]"
        b" /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    content = b"BT /F1 12 Tf 20 150 Td (Hello reader) Tj ET"
    objects.append(
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream"
    )
    data = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, obj in enumerate(objects, 1):
        offsets.append(len(data))
        data.extend(str(number).encode() + b" 0 obj\n" + obj + b"\nendobj\n")
    xref = len(data)
    data.extend(b"xref\n0 6\n0000000000 65535 f \n")
    for offset in offsets[1:]:
        data.extend(f"{offset:010} 00000 n \n".encode())
    data.extend(f"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode())
    result = read(bytes(data), "pdf")
    assert "Hello reader" in result.markdown
    assert result.pages[0].blocks[0].bbox[0] >= 19
    assert not any(w.code == "corrupted_page" for w in result.warnings)
