from __future__ import annotations

import io

from docx import Document
from docx.oxml import OxmlElement

from rag_pre_ingest.pipeline.readers import read_document, word_table_grid


def test_horizontal_merge_anchor_once_and_covered_cell_blank() -> None:
    table = Document().add_table(rows=2, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Merged heading"
    table.cell(0, 2).text = "Third"
    table.cell(1, 0).text = "Value"
    rows, cells = word_table_grid(table)
    assert rows == [["Merged heading", "", "Third"], ["Value", "", ""]]
    anchor = next(cell for cell in cells if cell.row == 0 and cell.col == 0)
    assert anchor.col_span == 2 and anchor.row_span == 1
    assert all(cell.is_header for cell in cells if cell.row == 0)
    assert not any(cell.row == 0 and cell.col == 1 for cell in cells)


def test_vertical_merge_anchor_span_and_no_repeated_text() -> None:
    table = Document().add_table(rows=3, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).merge(table.cell(2, 0)).text = "Spans rows"
    table.cell(1, 1).text = "One"
    table.cell(2, 1).text = "Two"
    rows, cells = word_table_grid(table)
    assert rows == [["Name", "Value"], ["Spans rows", "One"], ["", "Two"]]
    anchor = next(cell for cell in cells if cell.text == "Spans rows")
    assert anchor.row_span == 2 and anchor.col_span == 1
    assert not any(cell.row == 2 and cell.col == 0 for cell in cells)


def test_rectangle_merge_and_blank_slots() -> None:
    table = Document().add_table(rows=3, cols=3)
    table.cell(0, 0).merge(table.cell(1, 1)).text = "Rectangle"
    table.cell(0, 2).text = "Top"
    table.cell(1, 2).text = "Bottom"
    rows, cells = word_table_grid(table)
    assert rows == [["Rectangle", "", "Top"], ["", "", "Bottom"], ["", "", ""]]
    anchor = next(cell for cell in cells if cell.text == "Rectangle")
    assert anchor.row_span == 2 and anchor.col_span == 2
    assert sum(cell.text == "Rectangle" for cell in cells) == 1
    assert len(cells) == 6


def test_repeat_header_row_marked() -> None:
    table = Document().add_table(rows=2, cols=2)
    row_properties = table.rows[1]._tr.get_or_add_trPr()
    row_properties.append(OxmlElement("w:tblHeader"))
    _, cells = word_table_grid(table)
    assert all(cell.is_header for cell in cells)


def test_actual_word_result_markdown_and_spans() -> None:
    document = Document()
    table = document.add_table(rows=3, cols=3)
    table.cell(0, 0).text = "Header"
    table.cell(1, 0).merge(table.cell(2, 1)).text = "Unique merged value"
    table.cell(1, 2).text = "10"
    table.cell(2, 2).text = "20"
    output = io.BytesIO()
    document.save(output)
    result = read_document(output.getvalue(), "docx", "merged.docx", "", "")
    block = next(block for block in result.pages[0].blocks if block.type == "table")
    assert block.rows == 3 and block.columns == 3
    assert result.stats.tables.with_merged_cells == 1
    assert result.markdown.count("Unique merged value") == 1
    assert "| Unique merged value |  | 10 |" in result.markdown
    assert "|  |  | 20 |" in result.markdown
    assert not any(w.code == "corrupted_block" for w in result.warnings)
