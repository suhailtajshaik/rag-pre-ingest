from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from docx.shared import Inches
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from PIL import Image

from rag_pre_ingest.pipeline.readers import read_document, render_page


def picture() -> bytes:
    output = io.BytesIO()
    Image.new("RGB", (100, 60), "red").save(output, format="PNG")
    return output.getvalue()


def word_with_picture() -> bytes:
    document = Document()
    document.add_paragraph("Before image")
    document.add_picture(io.BytesIO(picture()), width=Inches(2))
    document.add_paragraph("Figure 1: red rectangle")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def test_word_embedded_image_order_assets_and_preview() -> None:
    data = word_with_picture()
    assets: dict[str, bytes] = {}
    callbacks: list[bool] = []
    result = read_document(
        data,
        "docx",
        "image.docx",
        "",
        "",
        image_assets=assets,
        on_page=lambda page: callbacks.append(
            all(block.id in assets for block in page.blocks if block.type == "image")
        ),
    )
    assert [b.type for b in result.pages[0].blocks] == ["paragraph", "image", "paragraph"]
    block = result.pages[0].blocks[1]
    assert block.id == "p1_image1"
    assert assets[block.id].startswith(b"\x89PNG")
    assert callbacks == [True]
    assert result.pages[0].blocks[2].bbox[1] > block.bbox[3]
    assert "image" not in result.markdown.replace("Before image", "")
    rendered = render_page(data, "docx", 1, result_page=result.pages[0], image_assets=assets)
    with Image.open(io.BytesIO(rendered)) as preview:
        assert preview.getpixel((int(block.bbox[0]) + 10, int(block.bbox[1]) + 10)) == (255, 0, 0)


def test_word_image_pixel_limit_becomes_warning() -> None:
    assets: dict[str, bytes] = {}
    result = read_document(
        word_with_picture(), "docx", "image.docx", "", "", image_assets=assets, max_image_pixels=50
    )
    assert not assets
    assert any(w.code == "unsupported_image" for w in result.warnings)
    assert "Before image" in result.markdown
    assert "Figure 1" in result.markdown


def test_excel_embedded_image() -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["A1"] = "Worksheet"
    sheet.add_image(ExcelImage(io.BytesIO(picture())), "B3")
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    assets: dict[str, bytes] = {}
    result = read_document(output.getvalue(), "xlsx", "image.xlsx", "", "", image_assets=assets)
    images = [block for block in result.pages[0].blocks if block.type == "image"]
    assert len(images) == 1
    assert images[0].id in assets
    assert result.stats.images.total == 1
    assert render_page(output.getvalue(), "xlsx", 1, image_assets=assets).startswith(b"\x89PNG")


def test_synthetic_word_fixture_exposes_actual_invoice() -> None:
    path = Path(__file__).resolve().parents[1] / "fixtures/synthetic/18-structured.docx"
    assets: dict[str, bytes] = {}
    result = read_document(path.read_bytes(), "docx", path.name, "", "", image_assets=assets)
    assert result.stats.images.total == 1
    assert len(assets) == 1
    with Image.open(io.BytesIO(next(iter(assets.values())))) as image:
        assert image.size == (1000, 600)


def test_external_word_image_relationship_is_never_followed() -> None:
    from docx.opc.constants import RELATIONSHIP_TYPE

    document = Document(io.BytesIO(word_with_picture()))
    relationship_id = document.part.relate_to(
        "http://127.0.0.1:9/not-fetched.png", RELATIONSHIP_TYPE.IMAGE, is_external=True
    )
    blip = document.element.body.xpath(".//a:blip")[0]
    blip.set(
        "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed",
        relationship_id,
    )
    output = io.BytesIO()
    document.save(output)
    assets: dict[str, bytes] = {}
    result = read_document(output.getvalue(), "docx", "external.docx", "", "", image_assets=assets)
    assert not assets
    assert any(w.code == "external_image_ignored" for w in result.warnings)
    assert "Before image" in result.markdown
