"""Preserve PDF source characters and resolve effective Word run typography."""

import io
from pathlib import Path

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Pt

from rag_pre_ingest.models.document import Block
from rag_pre_ingest.pipeline.readers import (
    attach_pdf_list_markers,
    paragraph_font_metadata,
    read_document,
    rectangle_text,
)

ROOT = Path(__file__).parents[2]


def test_rectangle_selection_preserves_source_spacing_and_apostrophes() -> None:
    text = "Sagamore l’ouest"
    characters = [(letter, (i * 5, 0, i * 5 + 4, 10)) for i, letter in enumerate(text)]
    assert rectangle_text(characters, (0, 0, 100, 10)) == text
    assert rectangle_text(characters, (0, 0, 39, 10)) == "Sagamore"


def test_literal_pdf_marker_attaches_by_geometry_without_mutating_inputs() -> None:
    bullet = Block(id="bullet", type="paragraph", text="•", bbox=(10, 15, 13, 18))
    content = Block(id="item", type="paragraph", text="An item", bbox=(25, 10, 100, 25))
    unrelated = Block(id="other", type="paragraph", text="Other line", bbox=(25, 40, 100, 55))
    blocks = attach_pdf_list_markers([bullet, content, unrelated])
    assert blocks[0].type == "list"
    assert blocks[0].markdown == "- An item"
    assert blocks[1] is unrelated
    assert bullet.type == "paragraph"
    assert content.text == "An item"


def test_effective_docx_font_uses_run_then_style_and_majority_text() -> None:
    doc = Document()
    style = doc.styles.add_style("Visual title", WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = doc.styles["Normal"]
    style.font.size = Pt(14)
    style.font.bold = True
    title = doc.add_paragraph("Manual title", style=style)
    assert paragraph_font_metadata(title, doc) == (14, 700)
    title.runs[0].font.bold = False
    assert paragraph_font_metadata(title, doc) == (14, 400)
    body = doc.add_paragraph()
    normal = body.add_run("This paragraph is primarily ordinary body text. ")
    normal.font.size = Pt(11)
    normal.bold = False
    emphasis = body.add_run("bold")
    emphasis.bold = True
    emphasis.font.size = Pt(14)
    assert paragraph_font_metadata(body, doc) == (11, 400)


def test_word_reader_exposes_manually_formatted_normal_style_title() -> None:
    doc = Document()
    title = doc.add_paragraph("Cooking with butter")
    title.runs[0].font.size = Pt(14)
    title.runs[0].bold = True
    doc.add_paragraph("Ordinary body text.")
    stream = io.BytesIO()
    doc.save(stream)
    blocks = read_document(stream.getvalue(), "docx", "", "", "").pages[0].blocks
    assert blocks[0].font_size == 14
    assert blocks[0].font_weight == 700
    assert blocks[1].font_size == 11
    assert blocks[1].font_weight == 400


def test_public_language_pdf_preserves_unspaced_word_fragments() -> None:
    path = ROOT / "tests/fixtures/public/.cache/w3c-language.pdf"
    if not path.exists():
        pytest.skip("Optional downloaded public corpus is not present.")
    result = read_document(path.read_bytes(), "pdf", path.name, "", "")
    assert "Sagamore" in result.markdown
    assert "You better" in result.markdown
    assert "This road is also" in result.markdown
    assert "(l’ouest)" in result.markdown
    assert not any(w.code == "corrupted_block" for w in result.warnings)


def test_public_pdf_lists_preserve_all_six_explicit_items() -> None:
    path = ROOT / "tests/fixtures/public/.cache/w3c-lists.pdf"
    if not path.exists():
        pytest.skip("Optional downloaded public corpus is not present.")
    result = read_document(path.read_bytes(), "pdf", path.name, "", "")
    lists = [block for block in result.pages[0].blocks if block.type == "list"]
    assert len(lists) == 6
    assert lists[0].markdown == "- Rocky High School"
    assert lists[3].markdown == "1. Snow-shoeing"


def test_pdf_discretionary_hyphen_uses_semantic_marker_not_pdfium_sentinel() -> None:
    path = ROOT / "tests/fixtures/synthetic/11-hyphen-wrap.pdf"
    result = read_document(path.read_bytes(), "pdf", path.name, "", "")
    assert "\ufffe" not in result.markdown
    assert "\u00ad" in result.markdown
    assert "well-known" in result.markdown
