import io
import signal
import subprocess
import sys
import time
from pathlib import Path
from unittest.mock import patch

import pypdfium2 as pdfium
import pytest
from openpyxl import Workbook
from PIL import Image

from rag_pre_ingest.models.document import Page
from rag_pre_ingest.pipeline.readers import page_bbox, page_deadline, read_document, render_page
from rag_pre_ingest.pipeline.validation import ValidationConfig, ValidationError, validate_upload


def test_later_tiff_frame_cannot_exceed_pixel_limit() -> None:
    output = io.BytesIO()
    Image.new("RGB", (2, 2)).save(
        output, format="TIFF", save_all=True, append_images=[Image.new("RGB", (20, 20))]
    )
    with pytest.raises(ValidationError) as error:
        validate_upload(output.getvalue(), "frames.tiff", ValidationConfig(max_image_pixels=100))
    assert error.value.code == "image_too_large"


def test_xlsx_upload_sheet_limit() -> None:
    workbook = Workbook()
    workbook.create_sheet("second")
    output = io.BytesIO()
    workbook.save(output)
    with pytest.raises(ValidationError) as error:
        validate_upload(output.getvalue(), "sheets.xlsx", ValidationConfig(max_pages=1))
    assert error.value.code == "too_many_pages"
    assert validate_upload(output.getvalue(), "sheets.xlsx", ValidationConfig()).page_count == 2


@pytest.mark.parametrize("rotation", [0, 90, 180, 270])
def test_rotated_pdf_coordinates_match_preview(rotation: int) -> None:
    path = Path("tests/fixtures/synthetic/01-clean.pdf")
    output = io.BytesIO()
    with pdfium.PdfDocument(path) as document:
        page = document[0]
        page.set_rotation(rotation)
        page.close()
        document.save(output)
    result = read_document(output.getvalue(), "pdf", "rotated.pdf", "hash", "id")
    page_result = result.pages[0]
    assert page_result.blocks
    for block in page_result.blocks:
        left, top, right, bottom = block.bbox
        assert 0 <= left <= right <= page_result.width
        assert 0 <= top <= bottom <= page_result.height


@pytest.mark.parametrize(
    ("rotation", "expected"),
    [
        (0, (10, 40, 30, 80)),
        (90, (20, 10, 60, 30)),
        (180, (70, 20, 90, 60)),
        (270, (40, 70, 80, 90)),
    ],
)
def test_bbox_crop_origin(rotation: int, expected: tuple[int, int, int, int]) -> None:
    assert page_bbox((20, 40, 40, 80), (10, 20, 110, 120), rotation) == expected


def test_office_preview_reuses_extracted_page() -> None:
    page = Page(page_number=1, width=800, height=1000)
    with patch("rag_pre_ingest.pipeline.readers.read_document", side_effect=AssertionError):
        assert render_page(b"", "xlsx", 1, result_page=page).startswith(b"\x89PNG")


def test_nested_deadline_cannot_extend_outer_timer() -> None:
    with pytest.raises(TimeoutError), page_deadline(0.03), page_deadline(1):
        time.sleep(0.1)
    assert signal.getitimer(signal.ITIMER_REAL)[0] == 0


def test_deadline_restores_only_remaining_outer_time() -> None:
    with page_deadline(1):
        before = signal.getitimer(signal.ITIMER_REAL)[0]
        with page_deadline(0.5):
            time.sleep(0.02)
        after = signal.getitimer(signal.ITIMER_REAL)[0]
        assert after < before - 0.015


def test_nested_expiry_does_not_leak_alarm_after_cleanup() -> None:
    # Isolate signals: a leaked alarm previously terminated pytest itself, or
    # raised a second TimeoutError outside the protected block under contention.
    program = """
import signal
import time
from rag_pre_ingest.pipeline.readers import page_deadline
for _ in range(1000):
    try:
        with page_deadline(0.0001), page_deadline(1):
            time.sleep(0.0002)
    except TimeoutError:
        pass
    assert signal.getitimer(signal.ITIMER_REAL)[0] == 0
time.sleep(0.01)
"""
    process = subprocess.run(
        [sys.executable, "-c", program], capture_output=True, text=True, timeout=15
    )
    assert process.returncode == 0, process.stderr


@pytest.mark.parametrize("name", ["07-scan.pdf", "08-mixed.pdf", "09-caption.pdf"])
def test_image_pdf_keeps_page_dimensions_and_image_blocks(name: str) -> None:
    data = (Path("tests/fixtures/synthetic") / name).read_bytes()
    result = read_document(data, "pdf", name, "hash", "id")
    assert result.pages[0].width > 0
    assert result.pages[0].height > 0
    assert any(block.type == "image" for block in result.pages[0].blocks)
    assert not any(w.code == "corrupted_page" for w in result.warnings)
    if name != "07-scan.pdf":
        assert result.markdown.strip()


def test_broken_image_object_preserves_direct_text() -> None:
    name = "08-mixed.pdf"
    data = (Path("tests/fixtures/synthetic") / name).read_bytes()
    with patch.object(pdfium.PdfImage, "get_bounds", side_effect=RuntimeError("broken image")):
        result = read_document(data, "pdf", name, "hash", "id")
    assert result.markdown.strip()
    assert result.pages[0].width > 0
    assert any(w.code == "corrupted_block" for w in result.warnings)
    assert not any(w.code == "corrupted_page" for w in result.warnings)
