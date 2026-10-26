import io
import time
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient

from rag_pre_ingest.config import Settings
from rag_pre_ingest.main import create_app


def test_background_read_and_all_result_endpoints(tmp_path: Path) -> None:
    word = Document()
    word.add_heading("Example heading", level=1)
    word.add_paragraph("Readable document content.")
    output = io.BytesIO()
    word.save(output)
    with TestClient(create_app(Settings(data_dir=tmp_path))) as client:
        uploaded = client.post("/v1/documents", files={"file": ("test.docx", output.getvalue())})
        assert uploaded.status_code == 202
        job_id = uploaded.json()["job_id"]
        for _ in range(100):
            job = client.get(f"/v1/jobs/{job_id}").json()
            if job["status"] in {"done", "failed"}:
                break
            time.sleep(0.05)
        assert job["status"] == "done", job
        result = client.get(f"/v1/jobs/{job_id}/result")
        assert result.status_code == 200
        assert "Readable document content." in result.json()["markdown"]
        assert client.get(f"/v1/jobs/{job_id}/markdown").text == result.json()["markdown"]
        image = client.get(f"/v1/jobs/{job_id}/pages/1/image")
        assert image.status_code == 200
        assert image.content.startswith(b"\x89PNG")
        assert client.get(f"/v1/jobs/{job_id}/pages/999/image").status_code == 404
        assert client.get("/v1/jobs").json()[0]["job_id"] == job_id
        events = client.get(f"/v1/jobs/{job_id}/events")
        assert events.status_code == 200
        assert '"status":"done"' in events.text
        assert client.get("/v1/jobs/missing/events").status_code == 404


def test_not_ready_and_missing_results(tmp_path: Path) -> None:
    with TestClient(create_app(Settings(data_dir=tmp_path, worker_enabled=False))) as client:
        assert client.get("/v1/jobs/missing/result").status_code == 404
        assert client.get("/v1/jobs/missing/markdown").status_code == 404
        word = Document()
        word.add_paragraph("Waiting")
        output = io.BytesIO()
        word.save(output)
        job_id = client.post("/v1/documents", files={"file": ("a.docx", output.getvalue())}).json()[
            "job_id"
        ]
        assert client.get(f"/v1/jobs/{job_id}/result").status_code == 409
        assert client.get(f"/v1/jobs/{job_id}/markdown").status_code == 409
