"""Direct reading only. Office pages are explicit synthetic previews, not pagination."""

from __future__ import annotations

import ctypes
import io
import re
import signal
import threading
import time
import unicodedata
from collections import Counter
from collections.abc import Callable, Iterator
from contextlib import closing, contextmanager
from datetime import UTC, datetime
from statistics import median
from typing import Any, Literal

import pypdfium2 as pdfium
from docx import Document as WordDocument
from openpyxl import load_workbook
from PIL import Image, ImageDraw

from rag_pre_ingest.models.document import (
    Block,
    Cell,
    Document,
    DocumentResult,
    Page,
    PageCounts,
    Stats,
    Warning,
    WarningCounts,
)
from rag_pre_ingest.pipeline.spreadsheet import trim_empty_tails


@contextmanager
def page_deadline(seconds: float) -> Iterator[None]:
    """Bound a page on the worker main thread; restore the previous timer."""
    if threading.current_thread() is not threading.main_thread():
        yield
        return
    previous_handler = signal.getsignal(signal.SIGALRM)

    def expired(signum: int, frame: object) -> None:
        raise TimeoutError("Page time limit exceeded.")

    signal.signal(signal.SIGALRM, expired)
    previous_timer = signal.getitimer(signal.ITIMER_REAL)
    started = time.monotonic()
    effective = min(seconds, previous_timer[0]) if previous_timer[0] > 0 else seconds
    signal.setitimer(signal.ITIMER_REAL, effective)
    timed_out = False
    try:
        yield
    except TimeoutError:
        timed_out = True
        raise
    finally:
        # Cancel before restoring handlers. An inner timeout already reports an
        # expired outer deadline; rearming it for a microsecond schedules a second
        # SIGALRM that can arrive after the outer scope restores SIG_DFL.
        signal.setitimer(signal.ITIMER_REAL, 0)
        remaining = previous_timer[0] - (time.monotonic() - started)
        signal.signal(signal.SIGALRM, previous_handler)
        if previous_timer[0] and remaining > 0:
            signal.setitimer(signal.ITIMER_REAL, remaining, previous_timer[1])
        elif previous_timer[0] and not timed_out:
            raise TimeoutError("Page time limit exceeded.")


def normalize_text(text: str) -> str:
    return unicodedata.normalize("NFKC", text).replace("\x00", "").strip()


def render_table(rows: list[list[str]], *, header: bool = True) -> str:
    if not rows:
        return ""
    width = max(map(len, rows))
    escaped = [
        [cell.replace("|", "\\|").replace("\n", " ") for cell in row] + [""] * (width - len(row))
        for row in rows
    ]
    lines = ["| " + " | ".join(row) + " |" for row in escaped]
    if not header:
        lines.insert(0, "| " + " | ".join([""] * width) + " |")
    lines.insert(1, "| " + " | ".join(["---"] * width) + " |")
    return "\n".join(lines)


def _warning(page: int | None, code: str, message: str) -> Warning:
    return Warning(page_number=page, code=code, message=message)


def page_bbox(
    box: tuple[float, float, float, float],
    bounds: tuple[float, float, float, float],
    rotation: int,
) -> tuple[float, float, float, float]:
    """Map PDF user coordinates into the rendered page's top-left coordinates."""
    left, bottom, right, top = bounds
    x0, y0, x1, y1 = box
    points: list[tuple[float, float]] = []
    for x, y in ((x0, y0), (x0, y1), (x1, y0), (x1, y1)):
        if rotation == 90:
            points.append((y - bottom, x - left))
        elif rotation == 180:
            points.append((right - x, y - bottom))
        elif rotation == 270:
            points.append((top - y, right - x))
        else:
            points.append((x - left, top - y))
    return (
        min(point[0] for point in points),
        min(point[1] for point in points),
        max(point[0] for point in points),
        max(point[1] for point in points),
    )


def _notify_page(
    page: Page, callback: Callable[[Page], None] | None, warnings: list[Warning]
) -> None:
    """Dispatch after extraction without misclassifying callback errors as source corruption."""
    if callback is not None:
        try:
            callback(page)
        except Exception:
            warnings.append(
                _warning(
                    page.page_number,
                    "page_callback_failed",
                    "Page follow-up could not be scheduled.",
                )
            )


def _pdf_words(textpage: Any, source: Any, number: int) -> list[Block]:
    words: list[Block] = []
    chars: list[str] = []
    boxes: list[tuple[float, float, float, float]] = []
    sizes: list[float] = []
    weights: list[int] = []

    def finish() -> None:
        if not chars:
            return
        value = normalize_text("".join(chars))
        if value and boxes:
            box = (
                min(b[0] for b in boxes),
                min(b[1] for b in boxes),
                max(b[2] for b in boxes),
                max(b[3] for b in boxes),
            )
            words.append(
                Block(
                    id=f"p{number}_word{len(words) + 1}",
                    type="paragraph",
                    text=value,
                    markdown=value,
                    bbox=box,
                    font_size=float(median(sizes)) if sizes else None,
                    font_weight=int(median(weights)) if weights else None,
                )
            )
        chars.clear()
        boxes.clear()
        sizes.clear()
        weights.clear()

    for index in range(textpage.count_chars()):
        value = textpage.get_text_range(index, 1)
        if pdfium.raw.FPDFText_IsHyphen(textpage, index) == 1:
            value = "\u00ad"
        if not value or value.isspace():
            finish()
            continue
        box = page_bbox(textpage.get_charbox(index), source.get_bbox(), source.get_rotation())
        size = float(pdfium.raw.FPDFText_GetFontSize(textpage, index))
        if boxes:
            previous = boxes[-1]
            # Separate spatially distant characters even when PDF lacks a space glyph.
            if abs(box[3] - previous[3]) > max(size * 0.5, 2) or box[0] - previous[2] > max(
                size * 0.35, 2
            ):
                finish()
        chars.append(value)
        boxes.append(box)
        sizes.append(size)
        weight = int(pdfium.raw.FPDFText_GetFontWeight(textpage, index))
        if weight > 0:
            weights.append(weight)
    finish()
    return words


def _filled_rectangle_line(
    points: list[tuple[float, float]],
) -> tuple[float, float, float, float] | None:
    """A thin filled axis-aligned rectangle represents one rule, not four edges."""
    if len(points) == 5 and points[0] == points[-1]:
        points = points[:-1]
    if len(points) != 4:
        return None
    for first, second in zip(points, [*points[1:], points[0]], strict=True):
        if abs(first[0] - second[0]) > 0.1 and abs(first[1] - second[1]) > 0.1:
            return None
    x0, y0 = min(p[0] for p in points), min(p[1] for p in points)
    x1, y1 = max(p[0] for p in points), max(p[1] for p in points)
    if not all(abs(x - x0) < 0.1 or abs(x - x1) < 0.1 for x, _ in points):
        return None
    if not all(abs(y - y0) < 0.1 or abs(y - y1) < 0.1 for _, y in points):
        return None
    width, height = x1 - x0, y1 - y0
    thickness, length = min(width, height), max(width, height)
    if not 0 < thickness <= 3 or length < max(4, 4 * thickness):
        return None
    if width >= height:
        return x0, (y0 + y1) / 2, x1, (y0 + y1) / 2
    return (x0 + x1) / 2, y0, (x0 + x1) / 2, y1


def rectangle_text(
    characters: list[tuple[str, tuple[float, float, float, float]]],
    bounds: tuple[float, float, float, float],
) -> str:
    """Select source characters without PDFium bounded-text spacing synthesis."""
    left, bottom, right, top = bounds
    selected = {
        index
        for index, (value, box) in enumerate(characters)
        if value
        and not value.isspace()
        and left - 0.01 <= (box[0] + box[2]) / 2 <= right + 0.01
        and bottom - 0.01 <= (box[1] + box[3]) / 2 <= top + 0.01
    }
    if not selected:
        return ""
    output: list[str] = []
    skipped = False
    for index in range(min(selected), max(selected) + 1):
        value = characters[index][0]
        if index in selected or value.isspace():
            if skipped and output and not output[-1].isspace():
                output.append("\n")
            output.append(value)
            skipped = False
        else:
            skipped = True
    return normalize_text("".join(output))


def _rectangle_span(
    characters: list[tuple[str, tuple[float, float, float, float]]],
    bounds: tuple[float, float, float, float],
) -> tuple[int, int] | None:
    left, bottom, right, top = bounds
    indices = [
        index
        for index, (value, box) in enumerate(characters)
        if value
        and not value.isspace()
        and left - 0.01 <= (box[0] + box[2]) / 2 <= right + 0.01
        and bottom - 0.01 <= (box[1] + box[3]) / 2 <= top + 0.01
    ]
    return (min(indices), max(indices)) if indices else None


def _path_lines(obj: Any, source: Any) -> list[tuple[float, float, float, float]]:
    """Read stroked segments and thin filled rules in transformed page coordinates."""
    fill, stroke = ctypes.c_int(), ctypes.c_int()
    if not pdfium.raw.FPDFPath_GetDrawMode(obj, ctypes.byref(fill), ctypes.byref(stroke)):
        return []
    if not fill.value and not stroke.value:
        return []
    matrix = obj.get_matrix()
    lines: list[tuple[float, float, float, float]] = []
    subpath: list[tuple[float, float]] = []
    edges: list[tuple[tuple[float, float], tuple[float, float]]] = []
    curved = False

    def transformed(point: tuple[float, float]) -> tuple[float, float]:
        x, y = matrix.on_point(*point)
        box = page_bbox((x, y, x, y), source.get_bbox(), source.get_rotation())
        return box[0], box[1]

    def finish(closed: bool) -> None:
        nonlocal curved
        centerline = (
            _filled_rectangle_line(subpath) if fill.value and closed and not curved else None
        )
        if centerline is not None:
            lines.append(centerline)
        elif stroke.value:
            for first, second in edges:
                x0, y0, x1, y1 = (
                    min(first[0], second[0]),
                    min(first[1], second[1]),
                    max(first[0], second[0]),
                    max(first[1], second[1]),
                )
                if (abs(x1 - x0) < 0.1 or abs(y1 - y0) < 0.1) and max(x1 - x0, y1 - y0) > 1:
                    lines.append((x0, y0, x1, y1))
        subpath.clear()
        edges.clear()
        curved = False

    for index in range(pdfium.raw.FPDFPath_CountSegments(obj)):
        segment = pdfium.raw.FPDFPath_GetPathSegment(obj, index)
        x, y = ctypes.c_float(), ctypes.c_float()
        if not pdfium.raw.FPDFPathSegment_GetPoint(segment, ctypes.byref(x), ctypes.byref(y)):
            raise ValueError("Invalid PDF path point.")
        point = transformed((float(x.value), float(y.value)))
        kind = pdfium.raw.FPDFPathSegment_GetType(segment)
        if kind == pdfium.raw.FPDF_SEGMENT_MOVETO:
            finish(False)
            subpath.append(point)
        elif kind == pdfium.raw.FPDF_SEGMENT_LINETO:
            if subpath:
                edges.append((subpath[-1], point))
            subpath.append(point)
        else:
            curved = True
            subpath.append(point)
        if pdfium.raw.FPDFPathSegment_GetClose(segment) and subpath:
            edges.append((subpath[-1], subpath[0]))
            finish(True)
    finish(False)
    return lines


def attach_pdf_list_markers(blocks: list[Block]) -> list[Block]:
    """Attach a literal detached PDF list marker to its adjacent source text."""
    output: list[Block] = []
    for block in blocks:
        if output:
            previous = output[-1]
            marker = previous.text.strip()
            height = block.bbox[3] - block.bbox[1]
            center = (previous.bbox[1] + previous.bbox[3]) / 2
            if (
                re.fullmatch(r"(?:[•▪◦●∙‣⁃]|[0-9]+[.)])", marker)
                and block.bbox[1] <= center <= block.bbox[3]
                and 0 <= block.bbox[0] - previous.bbox[2] <= height * 2
            ):
                text = marker + " " + block.text
                output[-1] = block.model_copy(
                    update={
                        "id": previous.id,
                        "type": "list",
                        "text": text,
                        "markdown": ("- " if marker in "•▪◦●∙‣⁃" else marker + " ") + block.text,
                        "bbox": (
                            previous.bbox[0],
                            min(previous.bbox[1], block.bbox[1]),
                            block.bbox[2],
                            max(previous.bbox[3], block.bbox[3]),
                        ),
                    }
                )
                continue
        output.append(block)
    return output


def _pdf(
    data: bytes,
    warnings: list[Warning],
    timeout: float,
    max_pages: int,
    on_page: Callable[[Page], None] | None = None,
) -> list[Page]:
    pages: list[Page] = []
    with pdfium.PdfDocument(data) as document:
        for index in range(min(len(document), max_pages)):
            number = index + 1
            try:
                with page_deadline(timeout), closing(document[index]) as source:
                    width, height = source.get_size()
                    blocks: list[Block] = []
                    words: list[Block] = []
                    ruling_lines: list[tuple[float, float, float, float]] = []
                    with closing(source.get_textpage()) as textpage:
                        text = normalize_text(textpage.get_text_bounded().replace("\x02", "\u00ad"))
                        characters = [
                            (
                                "\u00ad"
                                if pdfium.raw.FPDFText_IsHyphen(textpage, index) == 1
                                else textpage.get_text_range(index, 1).replace("\x02", "\u00ad"),
                                textpage.get_charbox(index),
                            )
                            for index in range(textpage.count_chars())
                        ]
                        try:
                            words = _pdf_words(textpage, source, number)
                        except TimeoutError:
                            raise
                        except Exception:
                            warnings.append(
                                _warning(
                                    number, "corrupted_block", "Word geometry could not be read."
                                )
                            )
                        # PDFium rectangles retain the source coordinates.
                        rectangles = textpage.count_rects()
                        previous_span: tuple[int, int] | None = None
                        for rect_index in range(rectangles):
                            left, bottom, right, top = textpage.get_rect(rect_index)
                            content = rectangle_text(characters, (left, bottom, right, top))
                            if content:
                                bounds = page_bbox(
                                    (left, bottom, right, top),
                                    source.get_bbox(),
                                    source.get_rotation(),
                                )
                                span = _rectangle_span(characters, (left, bottom, right, top))
                                if blocks and previous_span is not None and span is not None:
                                    previous = blocks[-1]
                                    fragment_height = max(
                                        bounds[3] - bounds[1], previous.bbox[3] - previous.bbox[1]
                                    )
                                    overlap = min(bounds[3], previous.bbox[3]) - max(
                                        bounds[1], previous.bbox[1]
                                    )
                                    if (
                                        previous_span[1] + 1 == span[0]
                                        and -0.1
                                        <= bounds[0] - previous.bbox[2]
                                        <= fragment_height * 0.35
                                        and overlap > 0.25 * fragment_height
                                    ):
                                        previous.text += content
                                        previous.markdown = previous.text
                                        previous.bbox = (
                                            min(bounds[0], previous.bbox[0]),
                                            min(bounds[1], previous.bbox[1]),
                                            max(bounds[2], previous.bbox[2]),
                                            max(bounds[3], previous.bbox[3]),
                                        )
                                        previous_span = (previous_span[0], span[1])
                                        continue
                                blocks.append(
                                    Block(
                                        id=f"p{number}_b{len(blocks) + 1}",
                                        type="paragraph",
                                        text=content,
                                        markdown=content,
                                        bbox=bounds,
                                    )
                                )
                                previous_span = span
                        if text and not blocks:
                            blocks.append(
                                Block(
                                    id=f"p{number}_b1",
                                    type="paragraph",
                                    text=text,
                                    markdown=text,
                                    bbox=(0, 0, width, height),
                                )
                            )
                    blocks = attach_pdf_list_markers(blocks)
                    for block in blocks:
                        relevant = [
                            word
                            for word in words
                            if block.bbox[0] - 1
                            <= (word.bbox[0] + word.bbox[2]) / 2
                            <= block.bbox[2] + 1
                            and block.bbox[1] - 1
                            <= (word.bbox[1] + word.bbox[3]) / 2
                            <= block.bbox[3] + 1
                        ]
                        sizes = [word.font_size for word in relevant if word.font_size is not None]
                        weights = [
                            word.font_weight for word in relevant if word.font_weight is not None
                        ]
                        block.font_size = float(median(sizes)) if sizes else None
                        block.font_weight = int(median(weights)) if weights else None
                    for obj in source.get_objects():
                        if obj.type == pdfium.raw.FPDF_PAGEOBJ_PATH:
                            try:
                                ruling_lines.extend(_path_lines(obj, source))
                            except TimeoutError:
                                raise
                            except Exception:
                                warnings.append(
                                    _warning(
                                        number,
                                        "corrupted_block",
                                        "A ruling path could not be read.",
                                    )
                                )
                        if isinstance(obj, pdfium.PdfImage):
                            try:
                                left, bottom, right, top = obj.get_bounds()
                                blocks.append(
                                    Block(
                                        id=f"p{number}_image{len(blocks) + 1}",
                                        type="image",
                                        bbox=page_bbox(
                                            (left, bottom, right, top),
                                            source.get_bbox(),
                                            source.get_rotation(),
                                        ),
                                        image_ref=f"pages/{number}/image",
                                    )
                                )
                            except TimeoutError:
                                raise
                            except Exception:
                                warnings.append(
                                    _warning(
                                        number,
                                        "corrupted_block",
                                        "An image block could not be read.",
                                    )
                                )
                    pages.append(
                        Page(
                            page_number=number,
                            width=width,
                            height=height,
                            blocks=blocks,
                            words=words,
                            ruling_lines=ruling_lines,
                        )
                    )
            except Exception as error:
                warnings.append(
                    _warning(
                        number,
                        "page_timeout" if isinstance(error, TimeoutError) else "corrupted_page",
                        "This PDF page could not be read.",
                    )
                )
                pages.append(Page(page_number=number, width=0, height=0))
            _notify_page(pages[-1], on_page, warnings)
    return pages


def _office_block(
    blocks: list[Block],
    text: str,
    kind: Literal["paragraph", "heading", "table", "list"] = "paragraph",
    level: int | None = None,
    cells: list[Cell] | None = None,
    rows: int | None = None,
    columns: int | None = None,
) -> None:
    y = blocks[-1].bbox[3] + 8 if blocks else 40.0
    blocks.append(
        Block(
            id=f"b{len(blocks) + 1}",
            type=kind,
            level=level,
            text=text,
            markdown=("#" * level + " " + text) if level else text,
            bbox=(40, y, 760, y + 40),
            cells=cells or [],
            rows=rows,
            columns=columns,
        )
    )


def _office_image(
    blocks: list[Block],
    data: bytes,
    assets: dict[str, bytes],
    page_number: int,
    max_pixels: int,
) -> None:
    """Decode a bounded embedded raster without loading external relationships."""
    with Image.open(io.BytesIO(data)) as image:
        if image.width * image.height > max_pixels:
            raise ValueError("Embedded image exceeds the pixel limit.")
        image.verify()
    with Image.open(io.BytesIO(data)) as image:
        image.load()
        scale = min(720 / image.width, 400 / image.height, 1)
        width, height = image.width * scale, image.height * scale
        output = io.BytesIO()
        image.convert("RGB").save(output, format="PNG")
    y = blocks[-1].bbox[3] + 8 if blocks else 40.0
    identifier = f"p{page_number}_image{sum(b.type == 'image' for b in blocks) + 1}"
    blocks.append(
        Block(
            id=identifier,
            type="image",
            bbox=(40, y, 40 + width, y + height),
            image_ref=f"images/{identifier}.png",
        )
    )
    assets[identifier] = output.getvalue()


def word_table_grid(table: Any) -> tuple[list[list[str]], list[Cell]]:
    """Read physical Word cells once; covered merge positions remain blank."""
    from docx.oxml.ns import qn
    from docx.table import _Cell

    width = len(table._tbl.tblGrid.gridCol_lst)
    rows: list[list[str]] = []
    cells: list[Cell] = []
    active: dict[int, Cell] = {}
    for row_number, row in enumerate(table._tbl.tr_lst):
        values = [""] * width
        properties = row.find(qn("w:trPr"))
        before = properties.find(qn("w:gridBefore")) if properties is not None else None
        column = int(before.get(qn("w:val"), "0")) if before is not None else 0
        repeated_header = properties is not None and properties.find(qn("w:tblHeader")) is not None
        next_active: dict[int, Cell] = {}
        extended: set[int] = set()
        for physical in row.tc_lst:
            properties = physical.find(qn("w:tcPr"))
            span_element = properties.find(qn("w:gridSpan")) if properties is not None else None
            span = int(span_element.get(qn("w:val"), "1")) if span_element is not None else 1
            if span < 1:
                raise ValueError("Invalid Word table column span.")
            merge = properties.find(qn("w:vMerge")) if properties is not None else None
            continuation = merge is not None and merge.get(qn("w:val"), "continue") != "restart"
            if column + span > len(values):
                values.extend([""] * (column + span - len(values)))
            if continuation:
                anchor = active.get(column)
                if anchor is None or any(
                    active.get(c) is not anchor for c in range(column, column + span)
                ):
                    raise ValueError("Word vertical merge has no matching anchor.")
                if id(anchor) not in extended:
                    anchor.row_span += 1
                    extended.add(id(anchor))
                for c in range(column, column + span):
                    next_active[c] = anchor
            else:
                value = normalize_text(_Cell(physical, table).text)
                anchor = Cell(
                    row=row_number,
                    col=column,
                    text=value,
                    is_header=row_number == 0 or repeated_header,
                    col_span=span,
                )
                cells.append(anchor)
                values[column] = value
                if merge is not None:
                    for c in range(column, column + span):
                        next_active[c] = anchor
            column += span
        rows.append(values)
        active = next_active
    width = max([width] + [len(row) for row in rows])
    return [row + [""] * (width - len(row)) for row in rows], cells


def word_list_marker(
    paragraph: Any, document: Any, counters: dict[tuple[str, int], int]
) -> tuple[str | None, int, bool]:
    """Resolve Word numbering without executing fields or inventing unknown formats."""
    from docx.oxml.ns import qn

    properties = paragraph._p.pPr
    numbering = properties.numPr if properties is not None else None
    style = paragraph.style
    while numbering is None and style is not None:
        properties = style.element.pPr
        numbering = properties.numPr if properties is not None else None
        style = style.base_style
    style_name = paragraph.style.name if paragraph.style else ""
    fallback = (
        "bullet"
        if style_name.startswith("List Bullet")
        else ("decimal" if style_name.startswith("List Number") else None)
    )
    suffix = style_name.rsplit(" ", 1)[-1]
    level = max(int(suffix) - 1, 0) if fallback and suffix.isdigit() else 0
    format_name: str | None
    if numbering is None:
        if fallback is None:
            return None, 0, False
        identifier, format_name, start = style_name, fallback, 1
    else:
        if numbering.numId is None or numbering.numId.val == 0:
            return None, 0, False
        identifier = str(numbering.numId.val)
        level = int(numbering.ilvl.val) if numbering.ilvl is not None else level
        root = document.part.numbering_part.element
        nums = [node for node in root.findall(qn("w:num")) if node.get(qn("w:numId")) == identifier]
        if not nums:
            return None, level, True
        abstract_id = nums[0].find(qn("w:abstractNumId"))
        abstracts = [
            node
            for node in root.findall(qn("w:abstractNum"))
            if abstract_id is not None
            and node.get(qn("w:abstractNumId")) == abstract_id.get(qn("w:val"))
        ]
        definition_level = int(numbering.ilvl.val) if numbering.ilvl is not None else 0
        definitions = [
            node
            for abstract in abstracts
            for node in abstract.findall(qn("w:lvl"))
            if node.get(qn("w:ilvl")) == str(definition_level)
        ]
        if not definitions:
            return None, level, True
        definition = definitions[0]
        format_element = definition.find(qn("w:numFmt"))
        format_name = format_element.get(qn("w:val")) if format_element is not None else None
        start_element = definition.find(qn("w:start"))
        start = int(start_element.get(qn("w:val"), "1")) if start_element is not None else 1
        for override in nums[0].findall(qn("w:lvlOverride")):
            if override.get(qn("w:ilvl")) == str(level):
                value = override.find(qn("w:startOverride"))
                if value is not None:
                    start = int(value.get(qn("w:val"), "1"))
    if format_name == "bullet":
        return "-", level, False
    if format_name != "decimal":
        return None, level, True
    key = (identifier, level)
    counters[key] = counters.get(key, start - 1) + 1
    for old_key in list(counters):
        if old_key[0] == identifier and old_key[1] > level:
            del counters[old_key]
    return str(counters[key]) + ".", level, False


def paragraph_font_metadata(paragraph: Any, document: Any) -> tuple[float | None, int | None]:
    """Resolve run, character style, paragraph style, and document font defaults."""
    from docx.oxml.ns import qn

    default_properties = document.styles.element.xpath("./w:docDefaults/w:rPrDefault/w:rPr")
    default_size: float | None = None
    default_bold = False
    if default_properties:
        size_element = default_properties[0].find(qn("w:sz"))
        bold_element = default_properties[0].find(qn("w:b"))
        if size_element is not None:
            default_size = float(size_element.get(qn("w:val"))) / 2
        if bold_element is not None:
            default_bold = bold_element.get(qn("w:val"), "true") not in {"0", "false", "off"}

    def styles(style: Any) -> list[Any]:
        fonts: list[Any] = []
        seen: set[str] = set()
        while style is not None and style.style_id not in seen:
            seen.add(style.style_id)
            fonts.append(style.font)
            style = style.base_style
        return fonts

    sizes: Counter[float] = Counter()
    weights: Counter[float] = Counter()
    for run in paragraph.runs:
        count = len(run.text.strip())
        if not count:
            continue
        fonts = [run.font, *styles(run.style), *styles(paragraph.style)]
        size = next((float(font.size.pt) for font in fonts if font.size is not None), default_size)
        bold = next((bool(font.bold) for font in fonts if font.bold is not None), default_bold)
        if size is not None:
            sizes[size] += count
        weights[700.0 if bold else 400.0] += count

    def dominant(values: Counter[float]) -> float | None:
        total = sum(values.values())
        count = 0
        for value, weight in sorted(values.items()):
            count += weight
            if count >= total / 2:
                return value
        return None

    size, weight = dominant(sizes), dominant(weights)
    return size, int(weight) if weight is not None else None


def _word(
    data: bytes,
    warnings: list[Warning],
    timeout: float,
    on_page: Callable[[Page], None] | None = None,
    image_assets: dict[str, bytes] | None = None,
    max_image_pixels: int = 40_000_000,
) -> list[Page]:
    assets = image_assets if image_assets is not None else {}
    document = WordDocument(io.BytesIO(data))
    blocks: list[Block] = []
    list_counters: dict[tuple[str, int], int] = {}
    # Iterate XML children to preserve paragraph/table order.
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for element in document.element.body:
        try:
            with page_deadline(timeout):
                if element.tag.endswith("}p"):
                    paragraph = Paragraph(element, document)
                    text = normalize_text(paragraph.text)
                    style = paragraph.style.name if paragraph.style else ""
                    level = (
                        int(style[-1])
                        if style and style.startswith("Heading ") and style[-1].isdigit()
                        else None
                    )
                    if text:
                        marker, indent, uncertain_list = word_list_marker(
                            paragraph, document, list_counters
                        )
                        if marker is not None:
                            listed = "  " * min(indent, 8) + marker + " " + text
                            _office_block(blocks, listed, "list")
                            box = blocks[-1].bbox
                            blocks[-1].bbox = (box[0] + min(indent, 8) * 24, box[1], box[2], box[3])
                        elif uncertain_list:
                            _office_block(blocks, text, "list")
                            warnings.append(
                                _warning(
                                    1,
                                    "lists_uncertain",
                                    "Unsupported Word numbering was preserved as plain list text.",
                                )
                            )
                        else:
                            _office_block(blocks, text, "heading" if level else "paragraph", level)
                        blocks[-1].font_size, blocks[-1].font_weight = paragraph_font_metadata(
                            paragraph, document
                        )
                    for blip in element.xpath(".//a:blip"):
                        try:
                            relationship_id = blip.get(
                                "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
                            )
                            if not relationship_id:
                                warnings.append(
                                    _warning(
                                        1,
                                        "external_image_ignored",
                                        "An external Word image was ignored.",
                                    )
                                )
                                continue
                            relationship = document.part.rels[relationship_id]
                            if relationship.is_external:
                                warnings.append(
                                    _warning(
                                        1,
                                        "external_image_ignored",
                                        "An external Word image was ignored.",
                                    )
                                )
                                continue
                            _office_image(
                                blocks, relationship.target_part.blob, assets, 1, max_image_pixels
                            )
                        except TimeoutError:
                            raise
                        except Exception:
                            warnings.append(
                                _warning(
                                    1,
                                    "unsupported_image",
                                    "An embedded Word image could not be decoded within limits.",
                                )
                            )
                elif element.tag.endswith("}tbl"):
                    table = Table(element, document)
                    rows, cells = word_table_grid(table)
                    _office_block(
                        blocks,
                        render_table(rows),
                        "table",
                        cells=cells,
                        rows=len(rows),
                        columns=max(map(len, rows), default=0),
                    )
        except Exception as error:
            warnings.append(
                _warning(
                    1,
                    "page_timeout" if isinstance(error, TimeoutError) else "corrupted_block",
                    "A Word block could not be read.",
                )
            )
    warnings.append(
        _warning(
            1,
            "synthetic_page",
            "Word preview uses a synthetic page; coordinates are not source pagination.",
        )
    )
    page = Page(
        page_number=1,
        width=800,
        height=max(1000, blocks[-1].bbox[3] + 60 if blocks else 1000),
        blocks=blocks,
    )
    pages = [page]
    _notify_page(page, on_page, warnings)
    return pages


def _excel(
    data: bytes,
    warnings: list[Warning],
    timeout: float,
    max_pages: int,
    on_page: Callable[[Page], None] | None = None,
    image_assets: dict[str, bytes] | None = None,
    max_image_pixels: int = 40_000_000,
) -> list[Page]:
    assets = image_assets if image_assets is not None else {}
    pages: list[Page] = []
    workbook = load_workbook(io.BytesIO(data), data_only=True, keep_links=False)
    try:
        for number, sheet in enumerate(workbook, 1):
            if number > max_pages:
                warnings.append(
                    _warning(None, "page_limit", "Additional worksheets exceed the page limit.")
                )
                break
            try:
                with page_deadline(timeout):
                    rows = [
                        [normalize_text(str(value)) if value is not None else "" for value in row]
                        for row in sheet.iter_rows(values_only=True)
                    ]
                    rows = trim_empty_tails(
                        rows,
                        [
                            (m.min_row - 1, m.min_col - 1, m.max_row - 1, m.max_col - 1)
                            for m in sheet.merged_cells.ranges
                        ],
                    )
                    cells = [
                        Cell(row=r, col=c, text=text)
                        for r, row in enumerate(rows)
                        for c, text in enumerate(row)
                    ]
                    for merged in sheet.merged_cells.ranges:
                        for cell in cells:
                            if cell.row == merged.min_row - 1 and cell.col == merged.min_col - 1:
                                cell.row_span = merged.max_row - merged.min_row + 1
                                cell.col_span = merged.max_col - merged.min_col + 1
                    blocks: list[Block] = []
                    _office_block(blocks, sheet.title, "heading", 1)
                    if any(any(row) for row in rows):
                        _office_block(
                            blocks,
                            render_table(rows),
                            "table",
                            cells=cells,
                            rows=len(rows),
                            columns=max(map(len, rows), default=0),
                        )
                    for image in sheet._images:
                        try:
                            _office_image(blocks, image._data(), assets, number, max_image_pixels)
                        except TimeoutError:
                            raise
                        except Exception:
                            warnings.append(
                                _warning(
                                    number,
                                    "unsupported_image",
                                    "A worksheet image could not be decoded within limits.",
                                )
                            )
                    pages.append(
                        Page(
                            page_number=number,
                            width=800,
                            height=max(
                                1000,
                                100 + len(rows) * 24,
                                blocks[-1].bbox[3] + 60 if blocks else 1000,
                            ),
                            blocks=blocks,
                        )
                    )
                    warnings.append(
                        _warning(
                            number,
                            "synthetic_page",
                            "Each worksheet is represented as a synthetic page.",
                        )
                    )
            except Exception as error:
                pages.append(Page(page_number=number, width=800, height=1000))
                warnings.append(
                    _warning(
                        number,
                        "page_timeout" if isinstance(error, TimeoutError) else "corrupted_page",
                        "This worksheet could not be read.",
                    )
                )
            _notify_page(pages[-1], on_page, warnings)
    finally:
        workbook.close()
    return pages


def read_document(
    data: bytes,
    file_type: str,
    file_name: str,
    file_hash: str,
    document_id: str,
    *,
    per_page_timeout_seconds: float = 30,
    max_pages: int = 500,
    on_page: Callable[[Page], None] | None = None,
    image_assets: dict[str, bytes] | None = None,
    max_image_pixels: int = 40_000_000,
) -> DocumentResult:
    started = time.monotonic()
    warnings: list[Warning] = []
    pages: list[Page] = []
    try:
        if file_type == "pdf":
            pages = _pdf(data, warnings, per_page_timeout_seconds, max_pages, on_page)
        elif file_type == "docx":
            pages = _word(
                data, warnings, per_page_timeout_seconds, on_page, image_assets, max_image_pixels
            )
        elif file_type == "xlsx":
            pages = _excel(
                data,
                warnings,
                per_page_timeout_seconds,
                max_pages,
                on_page,
                image_assets,
                max_image_pixels,
            )
        else:
            with Image.open(io.BytesIO(data)) as image:
                for index in range(min(getattr(image, "n_frames", 1), max_pages)):
                    try:
                        image.seek(index)
                        width, height = image.size
                        block = Block(
                            id=f"p{index + 1}_image1",
                            type="image",
                            bbox=(0, 0, width, height),
                            image_ref=f"pages/{index + 1}/image",
                        )
                        pages.append(
                            Page(page_number=index + 1, width=width, height=height, blocks=[block])
                        )
                    except Exception:
                        pages.append(Page(page_number=index + 1, width=0, height=0))
                        warnings.append(
                            _warning(
                                index + 1, "corrupted_page", "This image frame could not be read."
                            )
                        )
                    _notify_page(pages[-1], on_page, warnings)
    except Exception:
        warnings.append(_warning(None, "corrupted_page", "The document could not be read."))
    markdown = "\n\n".join(
        block.markdown
        for page in pages
        for block in page.blocks
        if block.type != "image" and block.markdown
    )
    if not markdown.strip():
        warnings.append(
            _warning(
                None,
                "empty_document",
                "No readable text was found; OCR is not available in this milestone.",
            )
        )
    blocks = [block for page in pages for block in page.blocks]
    stats = Stats(
        pages=PageCounts(total=len(pages), read=len(pages)),
        headings=sum(b.type == "heading" for b in blocks),
        paragraphs=sum(b.type == "paragraph" for b in blocks),
        warnings=WarningCounts(
            total=len(warnings), by_code=dict(Counter(w.code for w in warnings))
        ),
        duration_seconds=time.monotonic() - started,
    )
    stats.images.total = sum(b.type == "image" for b in blocks)
    stats.tables.total = sum(b.type == "table" for b in blocks)
    stats.tables.with_merged_cells = sum(
        any(c.row_span > 1 or c.col_span > 1 for c in b.cells) for b in blocks if b.type == "table"
    )
    return DocumentResult(
        document=Document(
            id=document_id,
            file_name=file_name,
            file_type=file_type,
            file_hash=file_hash,
            page_count=len(pages),
            processed_at=datetime.now(UTC),
        ),
        stats=stats,
        markdown=markdown,
        pages=pages,
        warnings=warnings,
    )


def render_page(
    data: bytes,
    file_type: str,
    page_number: int,
    max_pixels: int = 40_000_000,
    *,
    result_page: Page | None = None,
    image_assets: dict[str, bytes] | None = None,
) -> bytes:
    if page_number < 1:
        raise ValueError("Page numbers start at one.")
    output = io.BytesIO()
    if file_type == "pdf":
        with pdfium.PdfDocument(data) as document:
            with closing(document[page_number - 1]) as page:
                width, height = page.get_size()
                scale = min(1.5, (max_pixels / max(width * height, 1)) ** 0.5)
                bitmap = page.render(scale=scale)
                try:
                    bitmap.to_pil().save(output, format="PNG")
                finally:
                    bitmap.close()
    elif file_type in {"docx", "xlsx"}:
        page = result_page
        assets = image_assets if image_assets is not None else {}
        if page is None or (image_assets is None and any(b.type == "image" for b in page.blocks)):
            result = read_document(
                data,
                file_type,
                "preview",
                "",
                "preview",
                image_assets=assets,
                max_image_pixels=max_pixels,
            )
            if page is None:
                page = result.pages[page_number - 1]
        preview = Image.new("RGB", (int(page.width), min(int(page.height), 16000)), "white")
        draw = ImageDraw.Draw(preview)
        draw.text((20, 10), "Synthetic preview (not original pagination)", fill="gray")
        for block in page.blocks:
            if block.type == "image" and block.id in assets:
                with Image.open(io.BytesIO(assets[block.id])) as image:
                    size = (
                        max(1, int(block.bbox[2] - block.bbox[0])),
                        max(1, int(block.bbox[3] - block.bbox[1])),
                    )
                    preview.paste(
                        image.convert("RGB").resize(size), (int(block.bbox[0]), int(block.bbox[1]))
                    )
                continue
            draw.text((block.bbox[0], block.bbox[1]), block.text[:20000], fill="black")
        preview.save(output, format="PNG")
    else:
        with Image.open(io.BytesIO(data)) as image:
            image.seek(page_number - 1)
            image.convert("RGB").save(output, format="PNG")
    return output.getvalue()
