"""Bounded upload inspection; never extract text or follow external references."""

from __future__ import annotations

import io
import logging
import warnings
import zipfile
import zlib
from dataclasses import dataclass
from pathlib import PurePosixPath

import filetype
import pypdfium2 as pdfium
from docx import Document
from openpyxl import load_workbook
from PIL import Image, UnidentifiedImageError

logger = logging.getLogger(__name__)


@dataclass(frozen=True)
class ValidationConfig:
    max_bytes: int = 50 * 1024 * 1024
    max_pages: int = 500
    max_image_pixels: int = 40_000_000
    max_zip_ratio: float = 100.0
    max_zip_uncompressed_bytes: int = 200 * 1024 * 1024


@dataclass(frozen=True)
class ValidatedFile:
    file_type: str
    mime_type: str
    page_count: int | None
    extension_mismatch: bool


class ValidationError(ValueError):
    def __init__(self, code: str, message: str) -> None:
        self.code = code
        self.message = message
        super().__init__(message)


def validate_filename(filename: str) -> None:
    """Reject path-like client names independently of the host OS."""
    if (
        not filename
        or filename in {".", ".."}
        or any(character in filename for character in ("/", "\\", ":", "\x00"))
        or any(ord(character) < 32 for character in filename)
    ):
        raise ValidationError("unsafe_filename", "Provide a plain file name without a path.")


def validate_size(size: int, maximum: int) -> None:
    if size == 0:
        raise ValidationError("empty_file", "The uploaded file is empty.")
    if size > maximum:
        raise ValidationError("file_too_large", "The uploaded file exceeds the size limit.")


def _check_pages(count: int, config: ValidationConfig) -> None:
    if count > config.max_pages:
        raise ValidationError("too_many_pages", "The document exceeds the page limit.")
    if count == 0:
        raise ValidationError("corrupted_file", "The document has no pages.")


def inspect_archive(data: bytes, config: ValidationConfig) -> str:
    """Validate ZIP metadata before bounded CRC reads and OOXML parsing."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            names = [entry.filename for entry in entries]
            if len(names) != len(set(names)):
                raise ValidationError("corrupted_file", "The archive has duplicate entries.")
            total = 0
            for entry in entries:
                if entry.flag_bits & 1:
                    raise ValidationError("encrypted_file", "Encrypted files are not supported.")
                total += entry.file_size
                if total > config.max_zip_uncompressed_bytes or (
                    entry.file_size / max(entry.compress_size, 1) > config.max_zip_ratio
                ):
                    raise ValidationError(
                        "archive_limit_exceeded", "The archive exceeds expansion limits."
                    )
            if "[Content_Types].xml" not in names or "_rels/.rels" not in names:
                raise ValidationError(
                    "corrupted_file", "The Office archive is missing required parts."
                )
            if "word/document.xml" in names and "xl/workbook.xml" in names:
                raise ValidationError(
                    "corrupted_file", "The Office archive has ambiguous document parts."
                )
            if "word/document.xml" in names:
                kind = "docx"
            elif "xl/workbook.xml" in names:
                kind = "xlsx"
            else:
                raise ValidationError(
                    "unsupported_file_type", "Only DOCX and XLSX archives are supported."
                )
            if archive.testzip() is not None:
                raise ValidationError("corrupted_file", "The archive failed its integrity check.")
            return kind
    except ValidationError:
        raise
    except (
        zipfile.BadZipFile,
        zlib.error,
        RuntimeError,
        NotImplementedError,
        OSError,
        ValueError,
    ) as exc:
        raise ValidationError("corrupted_file", "The Office archive is damaged.") from exc


def _inspect_pdf(data: bytes, config: ValidationConfig) -> int:
    if not data.startswith(b"%PDF-"):
        raise ValidationError("corrupted_file", "The PDF header is invalid.")
    try:
        with pdfium.PdfDocument(data) as document:
            if pdfium.raw.FPDF_GetSecurityHandlerRevision(document) >= 0:
                raise ValidationError("encrypted_file", "Encrypted files are not supported.")
            count = len(document)
            _check_pages(count, config)
            return count
    except ValidationError:
        raise
    except pdfium.PdfiumError as exc:
        if getattr(exc, "err_code", None) == 4:
            raise ValidationError("encrypted_file", "Encrypted files are not supported.") from exc
        raise ValidationError("corrupted_file", "The PDF could not be opened.") from exc


def _inspect_image(data: bytes, config: ValidationConfig) -> int:
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("error", Image.DecompressionBombWarning)
            with Image.open(io.BytesIO(data)) as image:
                count = getattr(image, "n_frames", 1)
                _check_pages(count, config)
                for index in range(count):
                    try:
                        image.seek(index)
                    except (OSError, ValueError, EOFError):
                        # Later damaged frames are reported individually by the worker.
                        if index == 0:
                            raise
                        continue
                    if image.width * image.height > config.max_image_pixels:
                        raise ValidationError(
                            "image_too_large", "The image exceeds the pixel limit."
                        )
            with Image.open(io.BytesIO(data)) as checked:
                checked.verify()
            with Image.open(io.BytesIO(data)) as decoded:
                decoded.load()
            return int(count)
    except ValidationError:
        raise
    except (Image.DecompressionBombError, Image.DecompressionBombWarning) as exc:
        raise ValidationError("image_too_large", "The image exceeds the pixel limit.") from exc
    except (UnidentifiedImageError, OSError, ValueError, EOFError) as exc:
        raise ValidationError("corrupted_file", "The image could not be decoded.") from exc


def validate_upload(data: bytes, filename: str, config: ValidationConfig) -> ValidatedFile:
    """Validate first; callers may hash and persist only after this succeeds."""
    validate_filename(filename)
    validate_size(len(data), config.max_bytes)
    if data.startswith(bytes.fromhex("d0cf11e0a1b11ae1")):
        raise ValidationError(
            "encrypted_or_legacy_office", "Encrypted or legacy Office files are unsupported."
        )
    detected = filetype.guess(data)
    kind = detected.extension if detected is not None else None
    if data.startswith(b"PK"):
        kind = inspect_archive(data, config)
    if kind == "jpg":
        kind = "jpeg"
    elif kind == "tif":
        kind = "tiff"
    mime_types = {
        "pdf": "application/pdf",
        "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "png": "image/png",
        "jpeg": "image/jpeg",
        "gif": "image/gif",
        "webp": "image/webp",
        "tiff": "image/tiff",
        "bmp": "image/bmp",
    }
    if kind not in mime_types:
        raise ValidationError("unsupported_file_type", "The file type is not supported.")
    count: int | None = None
    if kind == "pdf":
        count = _inspect_pdf(data, config)
    elif kind in {"docx", "xlsx"}:
        try:
            if kind == "docx":
                Document(io.BytesIO(data))
            else:
                workbook = load_workbook(io.BytesIO(data), read_only=True, keep_links=False)
                try:
                    count = len(workbook.sheetnames)
                    _check_pages(count, config)
                finally:
                    workbook.close()
        except ValidationError:
            raise
        except Exception as exc:
            raise ValidationError(
                "corrupted_file", "The Office document could not be opened."
            ) from exc
    else:
        count = _inspect_image(data, config)
    extension = PurePosixPath(filename).suffix.lower().lstrip(".")
    aliases = {"jpg": "jpeg", "tif": "tiff"}
    mismatch = aliases.get(extension, extension) != kind
    if mismatch:
        logger.warning("Upload extension does not match detected type %s", kind)
    return ValidatedFile(kind, mime_types[kind], count, mismatch)
