"""Generate small reproducible scenario fixtures; never use private source documents."""

from __future__ import annotations

import io
import json
import re
import zipfile
import zlib
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.shared import Inches
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1] / "tests/fixtures/synthetic"


def pdf(pages: list[str], image: bytes | None = None) -> bytes:
    """Create a deterministic, valid PDF with a Helvetica font and optional RGB image."""
    objects: list[bytes] = [b"", b"", b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    image_id: int | None = None
    if image is not None:
        image_id = 4
        with Image.open(io.BytesIO(image)) as source:
            rgb = source.convert("RGB")
            content = zlib.compress(rgb.tobytes())
            objects.append(
                (
                    f"<< /Type /XObject /Subtype /Image /Width {rgb.width} /Height {rgb.height} "
                    f"/ColorSpace /DeviceRGB /BitsPerComponent 8 /Filter /FlateDecode "
                    f"/Length {len(content)} >>\nstream\n"
                ).encode()
                + content
                + b"\nendstream"
            )
    children = []
    for commands in pages:
        page_id = len(objects) + 1
        content_id = page_id + 1
        children.append(f"{page_id} 0 R")
        images = f"/XObject << /Im1 {image_id} 0 R >>" if image_id else ""
        objects.append(
            (
                f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                f"/Resources << /Font << /F1 3 0 R >> {images} >> /Contents {content_id} 0 R >>"
            ).encode()
        )
        stream = commands.encode("latin-1")
        objects.append(f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream")
    objects[0] = b"<< /Type /Catalog /Pages 2 0 R >>"
    objects[1] = f"<< /Type /Pages /Kids [{' '.join(children)}] /Count {len(pages)} >>".encode()
    data = bytearray(b"%PDF-1.4\n")
    offsets = []
    for index, obj in enumerate(objects, 1):
        offsets.append(len(data))
        data.extend(f"{index} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = len(data)
    data.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        data.extend(f"{offset:010} 00000 n \n".encode())
    data.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return bytes(data)


def text(value: str, x: int = 48, y: int = 700, size: int = 14) -> str:
    escaped = value.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    return f"BT /F1 {size} Tf {x} {y} Td ({escaped}) Tj ET\n"


def raster(
    value: str = "SCANNED DOCUMENT\nInvoice total 123.45",
    *,
    width: int = 1000,
    height: int = 600,
    format_name: str = "PNG",
) -> bytes:
    image = Image.new("RGB", (width, height), "white")
    candidates = [
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        Path("/System/Library/Fonts/Helvetica.ttc"),
    ]
    font = next(
        (ImageFont.truetype(str(path), 36) for path in candidates if path.exists()),
        ImageFont.load_default(size=36),
    )
    ImageDraw.Draw(image).multiline_text((50, 90), value, font=font, fill="black", spacing=25)
    output = io.BytesIO()
    image.save(output, format=format_name)
    return output.getvalue()


def canonical_zip(data: bytes) -> bytes:
    """Office producers embed wall-clock ZIP timestamps; canonicalize generated output."""
    output = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(data)) as source,
        zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target,
    ):
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, (2020, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            content = source.read(name)
            if name == "docProps/core.xml":
                content = re.sub(
                    rb"(<dcterms:modified[^>]*>)[^<]+", rb"\g<1>2020-01-01T00:00:00Z", content
                )
            target.writestr(info, content)
    return output.getvalue()


def generate(destination: Path = ROOT) -> list[dict[str, object]]:
    destination.mkdir(parents=True, exist_ok=True)
    manifest: list[dict[str, object]] = []

    def save(
        name: str,
        data: bytes,
        scenarios: list[str],
        expected: str = "accepted",
        config: dict[str, int] | None = None,
        expected_text: str = "",
    ) -> None:
        (destination / name).write_bytes(data)
        manifest.append(
            {
                "file": name,
                "scenarios": scenarios,
                "validation": expected,
                "config": config or {},
                "expected_text": expected_text,
            }
        )

    clean = pdf(
        [text("Clean document title", size=22) + text("Ordinary readable paragraph.", y=650)]
    )
    save(
        "01-clean.pdf",
        clean,
        ["direct_read", "heading"],
        "accepted",
        expected_text="Ordinary readable paragraph.",
    )
    repeated = [
        text("ACME report", y=760)
        + text(f"Unique content for section {n}", y=650)
        + text("Confidential", y=40)
        + text(str(n), x=300, y=20)
        for n in range(1, 4)
    ]
    save(
        "02-repeated.pdf",
        pdf(repeated),
        ["repeated_header", "repeated_footer", "page_numbers", "multipage"],
    )
    columns = (
        text("Full width heading", size=24)
        + text("Left first paragraph", x=40, y=640)
        + text("Left second paragraph", x=40, y=600)
        + text("Right first paragraph", x=330, y=640)
        + text("Right second paragraph", x=330, y=600)
    )
    save("03-two-columns.pdf", pdf([columns]), ["two_columns", "full_width_heading"])
    grid = "".join(f"{x} 480 m {x} 660 l S\n" for x in [40, 200, 360, 520]) + "".join(
        f"40 {y} m 520 {y} l S\n" for y in [480, 540, 600, 660]
    )
    values = (
        text("Name", 50, 625)
        + text("Qty", 210, 625)
        + text("Total", 370, 625)
        + text("Widget", 50, 565)
        + text("9.50", 370, 565)
        + text("Merged label", 50, 505)
    )
    save(
        "04-ruled-table.pdf", pdf([grid + values]), ["ruled_table", "blank_cells", "numeric_table"]
    )
    merged_grid = (
        "40 480 m 40 660 l S\n520 480 m 520 660 l S\n200 540 m 200 660 l S\n360 480 m 360 660 l S\n"
        + "".join(f"40 {y} m 520 {y} l S\n" for y in [480, 540, 600, 660])
    )
    save("05-merged-grid.pdf", pdf([merged_grid + values]), ["ruled_table", "merged_cells"])
    aligned = (
        text("Item", 40, 650)
        + text("Quantity", 230, 650)
        + text("Amount", 430, 650)
        + text("Bolts", 40, 610)
        + text("2", 230, 610)
        + text("10.00", 430, 610)
        + text("Nuts", 40, 570)
        + text("4", 230, 570)
        + text("5.00", 430, 570)
    )
    save("06-unruled-table.pdf", pdf([aligned]), ["unruled_table", "aligned_columns"])
    scan = raster()
    save(
        "07-scan.pdf",
        pdf(["q 550 0 0 330 30 350 cm /Im1 Do Q\n"], scan),
        ["scanned_pdf", "ocr_required"],
        "accepted",
        expected_text="Invoice total 123.45",
    )
    save(
        "08-mixed.pdf",
        pdf([text("Digital introduction") + "q 450 0 0 270 50 300 cm /Im1 Do Q\n"], scan),
        ["mixed_text_image", "region_ocr"],
    )
    save(
        "09-caption.pdf",
        pdf(
            ["q 400 0 0 240 50 400 cm /Im1 Do Q\n" + text("Figure 1: Invoice sample", x=50, y=375)],
            scan,
        ),
        ["caption", "image"],
    )
    logo = raster("ACME", width=220, height=150)
    save(
        "10-logo.pdf",
        pdf(
            [text(f"Body page {n}") + "q 60 0 0 40 500 740 cm /Im1 Do Q\n" for n in range(1, 4)],
            logo,
        ),
        ["repeated_logo", "decorative_image"],
    )
    save(
        "11-hyphen-wrap.pdf",
        pdf(
            [
                text("A broken para-", y=700)
                + text("graph continues here.", y=680)
                + text("A real well-known phrase.", y=640)
            ]
        ),
        ["line_wrap", "hyphen_repair", "real_hyphen"],
    )
    save("12-blank.pdf", pdf([""]), ["empty_document"])
    save("13-wrong-extension.docx", clean, ["extension_mismatch", "sniffed_pdf"])
    save("14-corrupted.pdf", b"%PDF-1.7\nnot a PDF", ["corrupted_input"], "corrupted_file")
    save("15-empty.pdf", b"", ["empty_input"], "empty_file")
    save("16-small-limit.pdf", clean, ["oversized_input"], "file_too_large", {"max_bytes": 20})
    save("17-page-limit.pdf", pdf(repeated), ["page_limit"], "too_many_pages", {"max_pages": 2})
    doc = Document()
    doc.core_properties.created = datetime(2020, 1, 1)
    doc.core_properties.modified = datetime(2020, 1, 1)
    doc.add_heading("Unicode and structure", 0)
    doc.add_heading("Chapter one", 1)
    doc.add_paragraph("Ligatures ﬁ ﬂ; nonbreaking\u00a0space; café résumé.")
    doc.add_paragraph("Curly “quotes” and apostrophe ’; en–dash em—dash; £ € ¥; ± × ÷ ≤ ≥ ≠ ∑.")
    doc.add_paragraph("العربية हिन्दी 中文 日本語 한국어")
    doc.add_paragraph("First bullet", style="List Bullet")
    doc.add_paragraph("First numbered item", style="List Number")
    table = doc.add_table(rows=3, cols=3)
    for c, value in enumerate(["Name", "Qty", "Total"]):
        table.cell(0, c).text = value
    table.cell(1, 0).text = "Widget"
    table.cell(1, 2).text = "12.50"
    table.cell(2, 0).merge(table.cell(2, 1)).text = "Merged label"
    doc.add_picture(io.BytesIO(scan), width=Inches(3))
    doc.add_paragraph("Figure 1: Embedded invoice")
    buffer = io.BytesIO()
    doc.save(buffer)
    save(
        "18-structured.docx",
        canonical_zip(buffer.getvalue()),
        [
            "unicode",
            "ligatures",
            "headings",
            "lists",
            "table",
            "blank_cells",
            "merged_cells",
            "embedded_image",
            "caption",
            "special_characters",
            "non_latin",
        ],
        "accepted",
        expected_text="Ligatures fi fl; nonbreaking space; café résumé.",
    )
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Numbers"
    sheet.append(["Item", "Qty", "Price", "Total"])
    sheet.append(["Widget", 2, None, 12.50])
    sheet.append(["Blank", None, 3.5, None])
    sheet["D4"] = "=SUM(D2:D3)"
    sheet.merge_cells("A5:B5")
    sheet["A5"] = "Merged total"
    sheet["A7"] = "café “quotes” £ € ¥ ± ≤ ≥"
    sheet["A8"] = "العربية हिन्दी 中文 日本語 한국어"
    book.create_sheet("Empty sheet")
    book.properties.created = datetime(2020, 1, 1)
    book.properties.modified = datetime(2020, 1, 1)
    buffer = io.BytesIO()
    book.save(buffer)
    book.close()
    save(
        "19-structured.xlsx",
        canonical_zip(buffer.getvalue()),
        [
            "numeric_cells",
            "blank_cells",
            "merged_cells",
            "formula",
            "multiple_sheets",
            "empty_sheet",
        ],
    )
    save(
        "20-scan.png",
        scan,
        ["standalone_image", "ocr_required"],
        "accepted",
        expected_text="Invoice total 123.45",
    )
    save("21-scan.jpg", raster(format_name="JPEG"), ["jpeg", "ocr_required"])
    buffer = io.BytesIO()
    with Image.open(io.BytesIO(scan)) as first:
        second = first.transpose(Image.Transpose.ROTATE_180)
        first.save(
            buffer, format="TIFF", compression="tiff_deflate", save_all=True, append_images=[second]
        )
    save("22-multipage.tiff", buffer.getvalue(), ["multipage_image", "rotated_scan"])
    save("23-image-limit.png", scan, ["pixel_limit"], "image_too_large", {"max_image_pixels": 100})
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("_rels/.rels", "<Relationships/>")
        archive.writestr("word/document.xml", "x" * 100000)
    save(
        "24-compression-limit.docx",
        canonical_zip(buffer.getvalue()),
        ["zip_expansion_limit"],
        "archive_limit_exceeded",
    )
    save(
        "25-encrypted-office.docx",
        bytes.fromhex("d0cf11e0a1b11ae1") + b"EncryptedPackage",
        ["encrypted_office_container"],
        "encrypted_or_legacy_office",
    )
    save("26-truncated.pdf", clean[:80], ["truncated_input"], "corrupted_file")
    save(
        "27-unsupported.txt",
        b"Plain text is outside upload allowlist.",
        ["unsupported_type"],
        "unsupported_file_type",
    )
    save("28-corrupt-image.png", scan[:40], ["corrupt_image"], "corrupted_file")
    save(
        "29-image-page-limit.tiff",
        (destination / "22-multipage.tiff").read_bytes(),
        ["image_page_limit"],
        "too_many_pages",
        {"max_pages": 1},
    )
    (destination / "manifest.json").write_text(
        json.dumps(
            {
                "description": "Synthetic scenarios only; real-world calibration remains required.",
                "fixtures": manifest,
            },
            indent=2,
        )
        + "\n"
    )
    return manifest


if __name__ == "__main__":
    print(f"Generated {len(generate())} synthetic fixtures in {ROOT}")
